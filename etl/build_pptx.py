"""On-site PPT: architecture first, then analysis."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from etl.config import ROOT

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1F, 0x7A, 0x8C)
GOLD = RGBColor(0xC4, 0x8B, 0x2F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT = RGBColor(0xF4, 0xF1, 0xEA)


def build_pptx(findings: dict, charts: dict | None = None) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _title(prs)
    _questions(prs)
    _architecture(prs)
    _elt_star(prs)
    _incremental_json(prs)
    charts = charts or {}
    _dq(prs, findings, charts)
    _sales(prs, findings, charts)
    _churn(prs, findings)
    _product(prs, findings, charts)
    _ops(prs, findings, charts)
    _tableau(prs, findings)
    _close(prs, findings)

    path = ROOT / "Retail_Sales_Analysis_Onsite_Deck.pptx"
    prs.save(path)
    return path


def _blank(prs):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def _bar(slide, color=NAVY):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _footer(slide, page: str):
    _text(slide, page, 12.2, 7.15, 1.0, 0.28, 11, GRAY, align=PP_ALIGN.RIGHT)
    _text(
        slide,
        "Retail Sales Case  |  Hybrid ELT  |  Snowflake mart  |  Tableau Public extracts",
        0.4,
        7.15,
        11,
        0.28,
        11,
        GRAY,
    )


def _text(slide, text, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = align
    return tf


def _bullets(slide, items, x, y, w, h, size=16):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.font.name = "Calibri"
        p.space_after = Pt(8)
        p.level = 0
    return tf


def _kpi_card(slide, x, y, title, value, subtitle=""):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.9), Inches(1.45)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.fill.background()
    _text(slide, title, x + 0.12, y + 0.08, 2.65, 0.32, 12, GRAY, bold=True)
    _text(slide, value, x + 0.12, y + 0.4, 2.65, 0.55, 22, NAVY, bold=True)
    if subtitle:
        _text(slide, subtitle, x + 0.12, y + 1.0, 2.65, 0.32, 11, TEAL)


def _title(prs):
    s = _blank(prs)
    bg = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()
    _text(s, "SENIOR DATA ANALYST  ·  CASE STUDY", 0.7, 1.6, 11, 0.4, 16, GOLD, bold=True)
    _text(s, "Retail Sales Performance &\nCustomer Retention", 0.7, 2.1, 12, 1.6, 36, WHITE, bold=True)
    _text(
        s,
        "Hybrid ELT  ·  Snowflake star-schema mart  ·  Tableau dashboard  ·  Daily incremental feeds",
        0.7,
        4.1,
        11,
        0.5,
        16,
        WHITE,
    )
    _text(
        s,
        "On-site presentation  |  Architecture, data quality, and decision-ready insights",
        0.7,
        6.5,
        11,
        0.4,
        14,
        GOLD,
    )


def _questions(prs):
    s = _blank(prs)
    _bar(s)
    _text(s, "What we are here to answer", 0.5, 0.4, 12, 0.5, 28, NAVY, bold=True)
    _bullets(
        s,
        [
            "Where is revenue coming from — store, region, hour, payment, member vs guest?",
            "Which members are at risk of churn, and what is a defensible churn definition?",
            "Which products, discounts, and baskets actually move the P&L?",
            "Can leadership trust the feed? (~56% of sales headers arrive without line items.)",
            "If this landed every day, how would production ingest, replay, and report on it?",
        ],
        0.6,
        1.2,
        12,
        5.5,
        18,
    )
    _footer(s, "02")


def _architecture(prs):
    s = _blank(prs)
    _bar(s)
    _text(s, "Architecture: hybrid ELT, not a one-off workbook", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    steps = [
        ("1. Landing", "landing/YYYY-MM-DD\nExcel + JSON drop"),
        ("2. Python", "Contract + DQ\nRAW load"),
        ("3. Snowflake", "VARIANT flatten\nStar schema mart"),
        ("4. Serve", "Tableau Desktop live\nPublic = extracts"),
    ]
    for i, (title, body) in enumerate(steps):
        x = 0.5 + i * 3.2
        card = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.3), Inches(2.9), Inches(2.4)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.fill.background()
        _text(s, title, x + 0.15, 1.45, 2.6, 0.45, 16, TEAL, bold=True)
        _text(s, body, x + 0.15, 2.05, 2.6, 1.4, 14, DARK)
    _bullets(
        s,
        [
            "RAW is the audit copy of the file that arrived today. Marts are rebuilt from RAW.",
            "Python parses files; Snowflake owns modeling, JSON flatten, and incremental MERGE-by-key.",
            "Tableau Public cannot connect to Snowflake — extracts are the publish path; live query is the production path.",
        ],
        0.6,
        4.0,
        12,
        2.6,
        16,
    )
    _footer(s, "03")


def _elt_star(prs):
    s = _blank(prs)
    _bar(s)
    _text(s, "Why ELT, and why a star schema", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    _text(s, "ELT (with a thin extract)", 0.5, 1.1, 6, 0.4, 18, TEAL, bold=True)
    _bullets(
        s,
        [
            "Replay: change a business rule without re-reading Excel.",
            "Audit: RAW proves what the daily file contained.",
            "JSON: VARIANT + LATERAL FLATTEN is native.",
            "Scale: transforms run in the warehouse.",
            "Python still owns file contracts: $ / pcs / %, Unnamed: 9.",
        ],
        0.5,
        1.55,
        6,
        4.5,
        15,
    )
    _text(s, "Star schema", 7.0, 1.1, 6, 0.4, 18, TEAL, bold=True)
    _bullets(
        s,
        [
            "fact_sales = one sale. fact_sales_items = one line.",
            "Dims: date, store (+ cashier), product, customer, payment.",
            "Guests: customer key is NULL. Never invent a customer.",
            "YTD / MTD / QTD / MoM need dim_date.",
            "Tableau filters on dims without blowing grain.",
        ],
        7.0,
        1.55,
        5.8,
        4.5,
        15,
    )
    _footer(s, "04")


def _incremental_json(prs):
    s = _blank(prs)
    _bar(s)
    _text(s, "Daily files and nested customers", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    _text(s, "Incremental landing", 0.5, 1.1, 6, 0.4, 18, TEAL, bold=True)
    _bullets(
        s,
        [
            "landing/2026-10-10/  then  landing/2026-10-11/",
            "Job flag: --as-of-date YYYY-MM-DD",
            "RAW stamped with _LOAD_DATE, _FILE_NAME, _INGESTED_AT",
            "MERGE on SALE_ID / ITEM_ID / CUSTOMER_ID; newer UPDATE_DATE wins",
            "Re-running the same folder is idempotent",
        ],
        0.5,
        1.55,
        6,
        4.8,
        15,
    )
    _text(s, "JSON → relational", 7.0, 1.1, 6, 0.4, 18, TEAL, bold=True)
    _bullets(
        s,
        [
            "RAW.CUSTOMERS.PAYLOAD as VARIANT",
            "LATERAL FLATTEN Phones → dim_customer_phone",
            "dim_customer keeps the Active phone for joins",
            "Tableau never queries nested JSON",
            "100 customers have an extra Inactive number — preserved, not dropped",
        ],
        7.0,
        1.55,
        5.8,
        4.8,
        15,
    )
    _footer(s, "05")


def _money(n: float) -> str:
    if abs(n) >= 1_000_000:
        return f" ${n/1_000_000:.2f}M"
    if abs(n) >= 1_000:
        return f" ${n:,.0f}"
    return f" ${n:,.2f}"


def _dq(prs, f, charts=None):
    s = _blank(prs)
    _bar(s)
    _text(s, "Data quality: treat the feed as a product", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    cov = f["kpi"]["item_coverage_pct"]
    _kpi_card(s, 0.5, 1.15, "ITEM COVERAGE", f"{cov:.1%}", "Headers with a basket")
    dirty = sum(int(r.get("row_count", 0)) for r in f.get("parse_issues", []))
    _kpi_card(s, 3.6, 1.15, "PRICE / QTY / %", f"{dirty:,}", "Dirty cells parsed in staging")
    _kpi_card(s, 6.7, 1.15, "GUEST KEYS", "NULL", "Never fabricate a customer")
    _kpi_card(s, 9.8, 1.15, "CASHIER", "1 : 1 store", "Not independent performance")
    _bullets(
        s,
        [
            "Unnamed: 9 is Customer ID — present for every Member, absent for every Guest.",
            "Line-item file only covers Sale_ID 100000–126666. Product KPIs are filtered with HAS_LINE_ITEMS.",
            "This is an operations finding: the items feed is lagging or truncated, not a reason to drop 33k sales.",
            "Gender = Unknown (569) is kept as its own segment.",
        ],
        0.6,
        2.75,
        12,
        1.55,
        13,
    )
    feed = (charts or {}).get("feed")
    if feed:
        s.shapes.add_picture(str(feed), Inches(0.45), Inches(4.35), Inches(12.3), Inches(2.55))
    _footer(s, "06")


def _sales(prs, f, charts=None):
    s = _blank(prs)
    _bar(s)
    _text(s, "Sales performance", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    k = f["kpi"]
    span = f"{f.get('sale_date_min', '')} – {f.get('sale_date_max', '')}".strip(" –")
    _kpi_card(s, 0.5, 1.1, "TRANSACTIONS", f"{k['txn_count']:,}", f"{k['stores']} stores  ·  {span}")
    _kpi_card(s, 3.6, 1.1, "REVENUE (WITH ITEMS)", _money(k["revenue_with_items"]).strip(), "Headers without lines excluded")
    _kpi_card(s, 6.7, 1.1, "AVG BASKET", _money(k["avg_basket_value"]).strip(), f"{k['avg_basket_size']:.1f} units")
    _kpi_card(
        s,
        9.8,
        1.1,
        "MEMBER / GUEST",
        f"{k['member_txns']:,} / {k['guest_txns']:,}",
        "Txn mix",
    )
    _text(s, "Store = cashier in this extract. Rank stores, not people.", 0.5, 2.7, 12, 0.3, 14, TEAL, bold=True)
    picture = (charts or {}).get("stores")
    if picture:
        s.shapes.add_picture(str(picture), Inches(0.45), Inches(3.05), Inches(12.3), Inches(3.9))
    _footer(s, "07")


def _churn(prs, f):
    s = _blank(prs)
    _bar(s)
    _text(s, "Member churn (definition first)", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    c = f["churn"]
    _kpi_card(s, 0.5, 1.15, "MEMBERS", f"{c['members']:,}", "Joined from JSON")
    _kpi_card(s, 3.6, 1.15, "CHURNED 90D", f"{c['churned_90d']:,}", f"{c['churn_rate']:.1%} of members")
    _kpi_card(s, 6.7, 1.15, "AVG FREQUENCY", f"{c['avg_frequency']:.1f}", "Purchases per member")
    _kpi_card(s, 9.8, 1.15, "AVG RECENCY", f"{c['avg_recency']:.0f}d", "Since last purchase")
    _bullets(
        s,
        [
            "Definition: Member with no purchase in the 90 days before the snapshot (max sale date). Guests cannot churn.",
            "Members buy about 2.4 times in 18 months, so a flat 90-day rule flags most of them. Win back the high-monetary band only.",
            "Dollar value is item revenue. Tickets with no lines add a visit but $0, so CLV is understated until the items feed is fixed.",
            "RFM quintiles are in mart_customer_rfm for the dashboard.",
        ],
        0.6,
        2.9,
        12,
        3.8,
        16,
    )
    _footer(s, "08")


def _product(prs, f, charts=None):
    s = _blank(prs)
    _bar(s)
    _text(s, "Assortment, discount, frequently bought together", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    picture = (charts or {}).get("products")
    if picture:
        s.shapes.add_picture(str(picture), Inches(0.35), Inches(1.05), Inches(8.1), Inches(5.7))
    pairs = [
        f"{r['PRODUCT_A_NAME']} + {r['PRODUCT_B_NAME']}  ·  lift {r['LIFT']:.2f}"
        for r in f["basket_pairs"][:5]
    ]
    _text(s, "Highest-lift pairs", 8.5, 1.15, 4.4, 0.4, 16, TEAL, bold=True)
    _bullets(s, pairs or ["No pairs"], 8.5, 1.7, 4.4, 5.0, 14)
    _footer(s, "09")


def _ops(prs, f, charts=None):
    s = _blank(prs)
    _bar(s)
    _text(s, "Time, payment, unusual activity", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    peak = f["peak_hour"]
    _kpi_card(
        s,
        0.5,
        1.15,
        "PEAK SLOT",
        f"{peak.get('SALE_DOW', '')} {int(peak.get('SALE_HOUR', 0)):02d}:00",
        f"{int(peak.get('TXN_COUNT', 0)):,} txns",
    )
    _kpi_card(s, 3.6, 1.15, "UNUSUAL TXNS", f"{f['unusual_count']:,}", "|z|≥3 or high-discount bulk")
    last = f["last_month"]
    mom = last.get("mom_growth_pct")
    mom_s = "n/a" if mom is None else f"{mom:.1%}"
    _kpi_card(s, 6.7, 1.15, "LAST MONTH MoM", mom_s, f"{last['year']}-{last['month']:02d}")
    _kpi_card(s, 9.8, 1.15, "PAYMENTS", "3-way split", "Cash / Debit / Credit roughly even")
    picture = (charts or {}).get("hourly")
    if picture:
        s.shapes.add_picture(str(picture), Inches(0.45), Inches(2.8), Inches(12.3), Inches(4.1))
    _footer(s, "10")


def _tableau(prs, f):
    s = _blank(prs)
    _bar(s)
    _text(s, "Dashboard and what you will demo", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    _bullets(
        s,
        [
            "Open extracts/tableau_public_pack.xlsx in Tableau Public Desktop (cannot live-query Snowflake).",
            "Pages: Executive KPI · Store & hour · Members & churn · Product & basket · Feed completeness.",
            f"Publish to {f.get('tableau_public', 'Tableau Public')}.",
            "If a licensed Tableau Desktop is available later, point it at RETAIL_SALES.MART — same grain, no remodel.",
            "Filters to keep: date, region, member/guest, HAS_LINE_ITEMS for product views.",
        ],
        0.6,
        1.2,
        12,
        5.4,
        17,
    )
    _footer(s, "11")


def _close(prs, f):
    s = _blank(prs)
    _bar(s, GOLD)
    _text(s, "Recommendations", 0.5, 0.4, 12, 0.5, 26, NAVY, bold=True)
    _bullets(
        s,
        [
            "Fix the items feed SLA — 55%+ of sales cannot support product, discount, or basket analytics.",
            "Run a 90-day win-back on churned members in the high monetary RFM band.",
            "Treat cashier KPIs as store KPIs until the grain changes.",
            "Keep RAW forever; rebuild marts. That is how you survive a logic change the night before a client review.",
            "Production path: Snowflake live → Tableau Desktop. Public extracts are the interview publish channel.",
        ],
        0.6,
        1.2,
        12,
        5.4,
        17,
    )
    _footer(s, "12")
